import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

df = pd.read_csv('ecommerce_sales_dataset.csv')
df['Order_Date'] = pd.to_datetime(df['Order_Date'])

# Create a clean Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "EDA_RESULTS"

# Styling definitions
font_title = Font(name="Calibri", size=16, bold=True, color="1F497D")
font_section = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
fill_section = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

font_kpi_label = Font(name="Calibri", size=10, bold=True, color="595959")
font_kpi_val = Font(name="Calibri", size=14, bold=True, color="1F497D")
fill_kpi = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

font_data = Font(name="Calibri", size=10)
font_data_bold = Font(name="Calibri", size=10, bold=True)
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

current_row = 1

def write_title(text):
    global current_row
    ws.cell(row=current_row, column=1, value=text).font = font_title
    current_row += 2

def write_section(title):
    global current_row
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    c = ws.cell(row=current_row, column=1, value=title)
    c.font = font_section
    c.fill = fill_section
    c.alignment = align_left
    current_row += 1

def write_table(dataframe, num_formats=None):
    global current_row
    # Headers
    for col_idx, col_name in enumerate(dataframe.columns, 1):
        c = ws.cell(row=current_row, column=col_idx, value=col_name)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = thin_border
    current_row += 1
    
    # Rows
    for _, row in dataframe.iterrows():
        for col_idx, val in enumerate(row, 1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = font_data
            c.border = thin_border
            col_name = dataframe.columns[col_idx - 1]
            
            if isinstance(val, (int, float)):
                c.alignment = align_right
                if num_formats and col_name in num_formats:
                    c.number_format = num_formats[col_name]
                elif "Margin" in col_name or "%" in col_name or "Share" in col_name or "Rate" in col_name:
                    c.number_format = "0.00%"
                elif "Revenue" in col_name or "Cost" in col_name or "Profit" in col_name or "AOV" in col_name or "Dollars" in col_name or "Freight" in col_name:
                    c.number_format = "$#,##0.00"
                elif isinstance(val, int) or "Orders" in col_name or "Quantity" in col_name or "Units" in col_name:
                    c.number_format = "#,##0"
                else:
                    c.number_format = "0.0000"
            else:
                c.alignment = align_left
        current_row += 1
    current_row += 1

write_title("OMNIVIZ SALES DATASET: EXPLORATORY DATA ANALYSIS (EDA) RESULTS")

# 1. Overall Business Performance
write_section("1. OVERALL BUSINESS PERFORMANCE (EXECUTIVE METRICS)")
tot_rev = df['Revenue'].sum()
tot_cost = df['Cost'].sum()
tot_profit = df['Profit'].sum()
tot_margin = tot_profit / tot_rev
mean_margin = df['Profit_Margin_%'].mean() / 100
tot_orders = len(df)
tot_qty = df['Quantity'].sum()
aov = df['Revenue'].mean()

kpi_data = pd.DataFrame([
    ["Total Revenue", tot_rev],
    ["Total Cost (COGS)", tot_cost],
    ["Total Gross Profit", tot_profit],
    ["Overall Weighted Gross Margin %", tot_margin],
    ["Average Order Profit Margin %", mean_margin],
    ["Total Orders", tot_orders],
    ["Total Quantity Sold", tot_qty],
    ["Average Order Value (AOV)", aov]
], columns=["Metric", "Exact Value"])
write_table(kpi_data)

# 2. Time Analysis
write_section("2. TIME ANALYSIS: YEARLY, QUARTERLY, AND SEASONAL PERFORMANCE")
year_df = df.groupby('Year').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    AOV=('Revenue', 'mean')
).reset_index()
year_df['Margin_%'] = year_df['Profit'] / year_df['Revenue']
write_table(year_df)

quarter_df = df.groupby('Quarter').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    AOV=('Revenue', 'mean')
).reset_index()
quarter_df['Margin_%'] = quarter_df['Profit'] / quarter_df['Revenue']
write_table(quarter_df)

season_df = df.groupby('Season').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    AOV=('Revenue', 'mean')
).reset_index()
season_df['Margin_%'] = season_df['Profit'] / season_df['Revenue']
write_table(season_df)

month_agg = df.groupby('Month').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
month_agg['Margin_%'] = month_agg['Profit'] / month_agg['Revenue']
month_agg['Rev_Share_%'] = month_agg['Revenue'] / tot_rev
write_table(month_agg)

# 3. Product Analysis
write_section("3. PRODUCT ANALYSIS: CATEGORIES, SUB-CATEGORIES, TOP & BOTTOM PRODUCTS")
cat_df = df.groupby('Category').agg(
    Orders=('Order_ID', 'count'),
    Units=('Quantity', 'sum'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
cat_df['Rev_Share_%'] = cat_df['Revenue'] / tot_rev
cat_df['Profit_Share_%'] = cat_df['Profit'] / tot_profit
cat_df['Margin_%'] = cat_df['Profit'] / cat_df['Revenue']
cat_df = cat_df.sort_values(by='Revenue', ascending=False)
write_table(cat_df)

subcat_df = df.groupby(['Category', 'Sub_Category']).agg(
    Orders=('Order_ID', 'count'),
    Units=('Quantity', 'sum'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum')
).reset_index()
subcat_df['Rev_Share_%'] = subcat_df['Revenue'] / tot_rev
subcat_df['Profit_Share_%'] = subcat_df['Profit'] / tot_profit
subcat_df['Margin_%'] = subcat_df['Profit'] / subcat_df['Revenue']
subcat_df = subcat_df.sort_values(by='Revenue', ascending=False)
write_table(subcat_df)

prod_df = df.groupby(['Category', 'Sub_Category', 'Product_Name']).agg(
    Orders=('Order_ID', 'count'),
    Units=('Quantity', 'sum'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum())
).reset_index()
prod_df['Margin_%'] = prod_df['Profit'] / prod_df['Revenue']

top10_rev = prod_df.sort_values(by='Revenue', ascending=False).head(10)
write_table(top10_rev)

top10_prof = prod_df.sort_values(by='Profit', ascending=False).head(10)
write_table(top10_prof)

bot10_prof = prod_df.sort_values(by='Profit', ascending=True).head(10)
write_table(bot10_prof)

# 4. Customer Analysis
write_section("4. CUSTOMER ANALYSIS: SEGMENTATION & GENDER PERFORMANCE")
seg_df = df.groupby('Customer_Segment').agg(
    Orders=('Order_ID', 'count'),
    Unique_Cust=('Customer_ID', 'nunique'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum())
).reset_index()
seg_df['AOV'] = seg_df['Revenue'] / seg_df['Orders']
seg_df['Margin_%'] = seg_df['Profit'] / seg_df['Revenue']
seg_df['Orders_per_Cust'] = seg_df['Orders'] / seg_df['Unique_Cust']
seg_df = seg_df.sort_values(by='Revenue', ascending=False)
write_table(seg_df)

gender_df = df.groupby('Customer_Gender').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum())
).reset_index()
gender_df['AOV'] = gender_df['Revenue'] / gender_df['Orders']
gender_df['Margin_%'] = gender_df['Profit'] / gender_df['Revenue']
write_table(gender_df)

# 5. Geographic Analysis
write_section("5. GEOGRAPHIC ANALYSIS: REGIONAL & COUNTRY LEVEL BREAKDOWN")
region_df = df.groupby('Region').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Ship_Days=('Shipping_Days', 'mean'),
    Avg_Ship_Cost=('Shipping_Cost', 'mean'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum())
).reset_index()
region_df['Rev_Share_%'] = region_df['Revenue'] / tot_rev
region_df['Margin_%'] = region_df['Profit'] / region_df['Revenue']
region_df = region_df.sort_values(by='Revenue', ascending=False)
write_table(region_df)

country_df = df.groupby(['Region', 'Country']).agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Profit=('Profit', 'sum'),
    Avg_Ship_Days=('Shipping_Days', 'mean'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum())
).reset_index()
country_df['Margin_%'] = country_df['Profit'] / country_df['Revenue']

top_country_rev = country_df.sort_values(by='Revenue', ascending=False).head(10)
write_table(top_country_rev)

top_country_prof = country_df.sort_values(by='Profit', ascending=False).head(10)
write_table(top_country_prof)

bot_country_prof = country_df.sort_values(by='Profit', ascending=True).head(10)
write_table(bot_country_prof)

# 6. Discount Analysis
write_section("6. DISCOUNT SENSITIVITY & MARGIN EROSION ANALYSIS")
disc_df = df.groupby('Discount').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Cost=('Cost', 'sum'),
    Profit=('Profit', 'sum'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum()),
    Loss_Dollars=('Profit', lambda x: x[x < 0].sum())
).reset_index()
disc_df['Margin_%'] = disc_df['Profit'] / disc_df['Revenue']
disc_df['Loss_Rate_%'] = disc_df['Loss_Orders'] / disc_df['Orders']
write_table(disc_df)

# 7. Logistics Analysis
write_section("7. LOGISTICS ANALYSIS: SHIPPING METHODS & TRANSIT BOTTLENECK")
ship_df = df.groupby('Shipping_Method').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Gross_Profit=('Profit', 'sum'),
    Freight_Total=('Shipping_Cost', 'sum'),
    Avg_Freight=('Shipping_Cost', 'mean'),
    Min_Days=('Shipping_Days', 'min'),
    Median_Days=('Shipping_Days', 'median'),
    Mean_Days=('Shipping_Days', 'mean'),
    Max_Days=('Shipping_Days', 'max')
).reset_index()
ship_df['Net_Profit_After_Shipping'] = ship_df['Gross_Profit'] - ship_df['Freight_Total']
ship_df['Gross_Margin_%'] = ship_df['Gross_Profit'] / ship_df['Revenue']
ship_df['Net_Margin_%'] = ship_df['Net_Profit_After_Shipping'] / ship_df['Revenue']
write_table(ship_df)

# 8. Order Status Lifecycle
write_section("8. ORDER STATUS: DELIVERED, RETURNED, CANCELLED & PROCESSING")
status_df = df.groupby('Order_Status').agg(
    Orders=('Order_ID', 'count'),
    Revenue=('Revenue', 'sum'),
    Profit=('Profit', 'sum'),
    Freight=('Shipping_Cost', 'sum'),
    Avg_Days=('Shipping_Days', 'mean')
).reset_index()
status_df['Volume_%'] = status_df['Orders'] / tot_orders
status_df['Rev_Share_%'] = status_df['Revenue'] / tot_rev
status_df['Profit_Share_%'] = status_df['Profit'] / tot_profit
status_df['Margin_%'] = status_df['Profit'] / status_df['Revenue']
write_table(status_df)

# 9. Profitability & Loss Audit
write_section("9. PROFITABILITY AUDIT: LOSS ORDERS, VALUE DESTROYERS & RISK CONCENTRATION")
loss_summary = pd.DataFrame([
    ["Total Loss-Making Orders", (df['Profit'] < 0).sum()],
    ["Loss Order Rate %", (df['Profit'] < 0).sum() / len(df)],
    ["Total Loss Dollars ($)", df[df['Profit'] < 0]['Profit'].sum()],
    ["Net Business Gross Profit ($)", tot_profit]
], columns=["Metric", "Calculated Value"])
write_table(loss_summary)

cat_loss = df.groupby('Category').agg(
    Total_Orders=('Order_ID', 'count'),
    Loss_Orders=('Profit', lambda x: (x < 0).sum()),
    Total_Loss_Dollars=('Profit', lambda x: x[x < 0].sum()),
    Net_Profit=('Profit', 'sum')
).reset_index()
cat_loss['Loss_Rate_%'] = cat_loss['Loss_Orders'] / cat_loss['Total_Orders']
cat_loss = cat_loss.sort_values(by='Loss_Orders', ascending=False)
write_table(cat_loss)

prod_loss = df[df['Profit'] < 0].groupby('Product_Name').agg(
    Loss_Orders=('Order_ID', 'count'),
    Loss_Dollars=('Profit', 'sum'),
    Avg_Discount=('Discount', 'mean'),
    Min_Discount=('Discount', 'min')
).reset_index().sort_values(by='Loss_Dollars', ascending=True).head(10)
write_table(prod_loss)

# 10. Statistical Relationships
write_section("10. STATISTICAL RELATIONSHIPS & PEARSON CORRELATION MATRIX")
pairs = [
    ("Revenue vs Profit", df['Revenue'].corr(df['Profit']), "Strong Positive", "Higher revenue translates directly into larger gross profit dollars."),
    ("Discount vs Profit", df['Discount'].corr(df['Profit']), "Moderate Negative", "Higher percentage discounts reduce total dollar profit."),
    ("Discount vs Profit Margin %", df['Discount'].corr(df['Profit_Margin_%']), "Very Strong Negative", "Steep inverse correlation: discounts directly crush percentage margins."),
    ("Quantity vs Revenue", df['Quantity'].corr(df['Revenue']), "Moderate Positive", "Larger basket sizes expand order revenue."),
    ("Shipping Cost vs Profit", df['Shipping_Cost'].corr(df['Profit']), "Zero Linear", "Shipping tier cost is completely independent of product gross margin."),
    ("Shipping Days vs Profit", df['Shipping_Days'].corr(df['Profit']), "Zero Linear", "Delivery transit duration has no linear association with order profitability."),
    ("Revenue vs Cost (COGS)", df['Revenue'].corr(df['Cost']), "Extremely Strong Positive", "Cost of goods sold tracks revenue in lockstep."),
    ("Shipping Cost vs Shipping Days", df['Shipping_Cost'].corr(df['Shipping_Days']), "Zero Linear", "CRITICAL ANOMALY: Paying 12x higher shipping cost provides 0 transit speed gain.")
]
corr_df = pd.DataFrame(pairs, columns=["Relationship Pair", "Pearson r", "Strength / Direction", "Operational Meaning"])
write_table(corr_df)

# Auto-fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

# Save Excel
wb.save('EDA_RESULTS.xlsx')
print("Successfully generated and saved EDA_RESULTS.xlsx with sheet 'EDA_RESULTS'!")
